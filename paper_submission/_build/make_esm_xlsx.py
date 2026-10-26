"""Online Resources 2-4 for the JIIM submission: each BMC additional-file CSV becomes ESM_<n>.xlsx with an
"About" sheet (article title, journal, caption; no author details, for double-blind review) and the data.
Usage: python _build/make_esm_xlsx.py"""
import csv
import json
import os
import subprocess

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, '..', '4_BMC_MedicalImaging_Submission', 'additional_files')
OUT = os.path.join(HERE, '..', '5_JIIM_Submission', 'online_resources')

TEXT_COLUMNS = {'pid', 'figshare_patient_id', 'name', 'path'}   # identifiers stay text


def number(v):
    try:
        return float(v)
    except ValueError:
        return v


info = json.loads(subprocess.check_output(['node', '-e', "const c=require('./content_jiim.js');console.log(JSON.stringify({title:c.meta.title,res:c.onlineResources}))"], cwd=HERE))
for i, r in enumerate(info['res'][1:], start=2):
    wb = Workbook()
    about = wb.active
    about.title = 'About'
    for row in (('Online Resource', f'{i} ({r["file"]})'), ('Article', info['title']), ('Journal', 'Journal of Imaging Informatics in Medicine'),
                ('Caption', r['title']), ('Description', r['description']), ('Authors', 'Given on the title page (double-blind review)')):
        about.append(row)
        about.cell(about.max_row, 1).font = Font(bold=True)
    about.column_dimensions['A'].width, about.column_dimensions['B'].width = 18, 120
    data = wb.create_sheet('Data')
    with open(os.path.join(SRC, f'Additional_file_{i}.csv'), newline='', encoding='utf-8') as f:
        rows = csv.reader(f)
        header = next(rows)
        data.append(header)
        for row in rows:
            data.append([v if h in TEXT_COLUMNS else number(v) for h, v in zip(header, row)])
    for c in data[1]:
        c.font = Font(bold=True)
    data.freeze_panes = 'A2'
    wb.save(os.path.join(OUT, r['file']))
    print(r['file'], data.max_row - 1, 'rows')
